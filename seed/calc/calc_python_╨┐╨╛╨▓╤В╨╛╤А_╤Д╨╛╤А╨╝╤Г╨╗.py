import openpyxl,sys
def run(fn,lease=True):
    ws=openpyxl.load_workbook(fn).active
    out=[]
    for c in "FHJL":
        v=lambda r: ws[f"{c}{r}"].value or 0
        den=v(29)-v(28)-v(27)
        k1=(v(20)+v(37))/den; k2=(v(20)+v(19)-v(38)+v(18)-v(35))/den
        k3=(v(21)-v(35)-v(36)-v(38)-v(39))/den; k4=(v(25)+v(27)+v(28))/v(30)
        k5=v(11)/v(10); k6=v(12)/v(10)
        g1=3 if k1<.05 else 2 if k1<.1 else 1
        g2=3 if k2<.5 else 2 if k2<.8 else 1
        g3=3 if k3<1 else 2 if k3<1.5 else 1
        g4=(3 if k4<.15 else 2 if k4<.25 else 1) if lease else (3 if k4<.25 else 2 if k4<.4 else 1)
        g5=3 if k5<0 else 2 if k5<.1 else 1
        g6=3 if k6<0 else 2 if k6<.06 else 1
        S=.05*g1+.1*g2+.4*g3+.2*g4+.15*g5+.1*g6
        cl=1 if (k5>=.1 and S<=1.25) else 2 if (S<=2.35 and k5>=0) else 3
        out.append((ws[f"{c}9"].value.date(),[round(x,3) for x in (k1,k2,k3,k4,k5,k6)],(g1,g2,g3,g4,g5,g6),round(S,2),cl))
    return out
for fn in sys.argv[1:]:
    print(fn)
    for lease in (True,False):
        print(" blank-D5(trade/lease thresholds)" if lease else " D5 filled other industry")
        for r in run(fn,lease): print("  ",r)
